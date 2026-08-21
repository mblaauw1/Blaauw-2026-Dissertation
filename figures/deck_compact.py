"""Compact deck: same manifest as deck.py, but every image is downsampled to a
slide-appropriate JPEG. Outputs a much smaller .pptx AND a portable .pdf.
Full-res deck.py is unchanged; this is the lightweight companion."""
import os
from PIL import Image
# Some whole-movie timestrips legitimately exceed PIL's 178 Mpixel decompression-bomb guard
# (measured 200,475,438 px on 2026-07-29), which aborted this build. These are our own rendered
# figures, not untrusted input, so the guard is not protecting anything here.
Image.MAX_IMAGE_PIXELS = None
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

ROOT="/Volumes/4 MB/ablation_figures_20260625"
IMGCACHE=os.path.join(ROOT,"_compact_img"); os.makedirs(IMGCACHE,exist_ok=True)
MAXW=2200            # max width px for any embedded image (timestrips stay legible)
JPEGQ=82

# --- reuse the manifest from deck.py (exec only the GROUPS block) ---
ns={}
top=open(os.path.join(ROOT,"deck.py")).read().split("# ---------------- PPTX")[0]
exec(top,ns)
GROUPS=ns["GROUPS"]

def compact(img):
    """downsample to JPEG, return cached path. white-background flatten for transparency."""
    src=os.path.join(ROOT,img)
    if not os.path.isfile(src): return None
    out=os.path.join(IMGCACHE,img.replace("/","__").rsplit(".",1)[0]+".jpg")
    os.makedirs(os.path.dirname(out),exist_ok=True) if "/" in out else None
    if os.path.isfile(out) and os.path.getmtime(out)>=os.path.getmtime(src):
        return out
    im=Image.open(src)
    if im.mode in ("RGBA","P","LA"):
        bg=Image.new("RGB",im.size,(255,255,255)); bg.paste(im.convert("RGBA"),mask=im.convert("RGBA").split()[-1]); im=bg
    else: im=im.convert("RGB")
    if im.width>MAXW:
        im=im.resize((MAXW,max(1,round(im.height*MAXW/im.width))),Image.LANCZOS)
    im.save(out,"JPEG",quality=JPEGQ,optimize=True)
    return out

# ================= compact PPTX =================
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
blank=prs.slide_layouts[6]
def sep_slide(g):
    s=prs.slides.add_slide(blank)
    r,gr,b=int(g["color"][0:2],16),int(g["color"][2:4],16),int(g["color"][4:6],16)
    box=s.shapes.add_shape(1,0,0,prs.slide_width,prs.slide_height); box.fill.solid()
    box.fill.fore_color.rgb=RGBColor(r,gr,b); box.line.fill.background()
    tb=s.shapes.add_textbox(Inches(1),Inches(3),Inches(11.3),Inches(1.6)).text_frame; tb.word_wrap=True
    p=tb.paragraphs[0]; p.text=f"GROUP {g['n']}"; p.font.size=Pt(28); p.font.bold=True
    p.font.color.rgb=RGBColor(255,255,255); p.alignment=PP_ALIGN.CENTER
    p2=tb.add_paragraph(); p2.text=g["title"]; p2.font.size=Pt(20); p2.font.color.rgb=RGBColor(255,255,255); p2.alignment=PP_ALIGN.CENTER
def img_slide(path,cap):
    s=prs.slides.add_slide(blank)
    w,h=Image.open(path).size; ar=w/h
    maxw,maxh=Inches(12.3),Inches(6.3); iw=maxw; ih=Emu(int(iw/ar))
    if ih>maxh: ih=maxh; iw=Emu(int(ih*ar))
    s.shapes.add_picture(path,int((prs.slide_width-iw)/2),Inches(0.5),iw,ih)
    tb=s.shapes.add_textbox(Inches(0.5),Inches(6.95),Inches(12.3),Inches(0.5)).text_frame
    tb.paragraphs[0].text=cap; tb.paragraphs[0].font.size=Pt(12); tb.paragraphs[0].alignment=PP_ALIGN.CENTER
def ph_slide(txt):
    s=prs.slides.add_slide(blank)
    box=s.shapes.add_shape(1,Inches(1.5),Inches(2.5),Inches(10.3),Inches(2.5))
    box.fill.solid(); box.fill.fore_color.rgb=RGBColor(245,245,245); box.line.color.rgb=RGBColor(180,180,180); box.line.dash_style=2
    tf=box.text_frame; tf.word_wrap=True
    p=tf.paragraphs[0]; p.text="◻ PLACEHOLDER"; p.font.size=Pt(16); p.font.bold=True; p.font.color.rgb=RGBColor(120,120,120); p.alignment=PP_ALIGN.CENTER
    p2=tf.add_paragraph(); p2.text=txt; p2.font.size=Pt(13); p2.font.color.rgb=RGBColor(90,90,90); p2.alignment=PP_ALIGN.CENTER

for g in GROUPS:
    sep_slide(g)
    for sl in g["slides"]:
        cp=compact(sl["img"]) if "img" in sl else None
        if cp: img_slide(cp,sl["cap"])
        else: ph_slide(sl.get("ph") or sl.get("cap","(missing image)"))
pptx_out=os.path.join(ROOT,"ablation_figures_compact.pptx"); prs.save(pptx_out)

# ================= portable PDF (landscape, 13.333x7.5in @72pt/in) =================
PW,PH=13.333*72,7.5*72
pdf_out=os.path.join(ROOT,"ablation_figures.pdf")
c=canvas.Canvas(pdf_out,pagesize=(PW,PH))
def pdf_sep(g):
    r,gr,b=int(g["color"][0:2],16)/255,int(g["color"][2:4],16)/255,int(g["color"][4:6],16)/255
    c.setFillColorRGB(r,gr,b); c.rect(0,0,PW,PH,fill=1,stroke=0)
    c.setFillColorRGB(1,1,1)
    c.setFont("Helvetica-Bold",30); c.drawCentredString(PW/2,PH/2+18,f"GROUP {g['n']}")
    c.setFont("Helvetica",20); c.drawCentredString(PW/2,PH/2-18,g["title"])
    c.showPage()
def pdf_img(path,cap):
    c.setFillColorRGB(1,1,1); c.rect(0,0,PW,PH,fill=1,stroke=0)
    iw,ih=Image.open(path).size; ar=iw/ih
    maxw,maxh=PW-72,PH-90; dw=maxw; dh=dw/ar
    if dh>maxh: dh=maxh; dw=dh*ar
    c.drawImage(ImageReader(path),(PW-dw)/2,PH-46-dh,dw,dh,preserveAspectRatio=True,mask='auto')
    c.setFillColorRGB(0,0,0); c.setFont("Helvetica",10)
    c.drawCentredString(PW/2,16,cap[:180])
    c.showPage()
def pdf_ph(txt):
    c.setFillColorRGB(.96,.96,.96); c.rect(0,0,PW,PH,fill=1,stroke=0)
    c.setFillColorRGB(.4,.4,.4); c.setFont("Helvetica-Bold",16); c.drawCentredString(PW/2,PH/2+12,"◻ PLACEHOLDER")
    c.setFont("Helvetica",12); c.drawCentredString(PW/2,PH/2-12,txt[:160]); c.showPage()
for g in GROUPS:
    pdf_sep(g)
    for sl in g["slides"]:
        cp=compact(sl["img"]) if "img" in sl else None
        if cp: pdf_img(cp,sl["cap"])
        else: pdf_ph(sl.get("ph") or sl.get("cap",""))
c.save()

print("compact pptx:",round(os.path.getsize(pptx_out)/1e6,1),"MB ->",pptx_out)
print("pdf        :",round(os.path.getsize(pdf_out)/1e6,1),"MB ->",pdf_out)
