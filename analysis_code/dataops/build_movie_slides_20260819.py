#!/usr/bin/env python3
"""Three slide decks organising the processed movies behind the figures.

USER 2026-08-19 (to-do list 0819 1pm), the NON-FIGURE item — verbatim:
    "Organize (processed) movies relevant to the figures (70-something?) in some google slides docs: a slides
     for the movies relevant to the main figures, a slides for the movies relevant to the supplemental, and a
     third slides for the movies relevant to the other three ai files (but still separate into three groups by
     putting a title slide to lead off the group of movies relevant to each different ai file."

WHAT IS BUILT
    MOVIES_MAIN_FIGURES.pptx          movies whose figure sits on META_FIGURES_20260814
    MOVIES_SUPPLEMENTAL.pptx          movies whose figure sits on META_FIGURES_20260813_supplemental
    MOVIES_OTHER_THREE_FILES.pptx     the remaining three decks, each led off by its own TITLE SLIDE
                                      (NEW_FIGURES_20260804, other.ai, NEW_TIMESTRIPS_20260804)
    .pptx, because Google Slides imports PowerPoint directly (File > Import slides, or drop it in Drive and
    "Open with Google Slides") and nothing here needs the Slides API.

WHERE THE MOVIES AND THEIR ASSIGNMENTS COME FROM
    `3_MOVIES/timestrip_movies_20260817/`, built 2026-08-17 by `custom_deck_movies_20260817.py`, one clip per
    full timestrip on the decks, already verified playable and trimmed. Each file is named
        <DECK>_AB<artboard>__<figure>__<batch>__<role>.mp4
    so the deck, artboard, figure and cell are IN THE FILENAME -- the grouping is read from the movies
    themselves, not from a list that could go stale.

WHAT EACH MOVIE SLIDE CARRIES
    a poster frame pulled from the clip (so the deck is browsable without opening anything), the figure name,
    the deck and artboard it belongs to, the cell, the portion (ablation / monitoring), the duration, and the
    full path on the 4 MB drive. A slide is therefore enough to FIND the movie, which is what makes the deck
    useful next to the artboards.

    A slide cannot PLAY a local file in Google Slides -- Slides only plays Drive or YouTube video. The path is
    given so the clip can be opened from the drive, and if she later wants them playable in the deck the same
    script can be pointed at Drive file ids once the clips are uploaded.
"""
import collections
import glob
import os
import re
import subprocess
import sys

import cv2

ROOT = "/Volumes/4 MB"
MOV = f"{ROOT}/3_MOVIES/timestrip_movies_20260817"
OUTDIR = f"{ROOT}/3_MOVIES/movie_slides_20260819"
POSTER = f"{OUTDIR}/_posters"

DECK_TITLE = {
    "META_FIGURES_20260814": "MAIN FIGURES — META_FIGURES_20260814.ai",
    "META_FIGURES_20260813_supplemental": "SUPPLEMENTAL — META_FIGURES_20260813_supplemental.ai",
    "NEW_FIGURES_20260804": "NEW_FIGURES_20260804.ai",
    "other": "other.ai  (was 'supplemental.ai', renamed 2026-08-19)",
    "NEW_TIMESTRIPS_20260804": "NEW_TIMESTRIPS_20260804.ai",
}
GROUPS = [
    ("MOVIES_MAIN_FIGURES", ["META_FIGURES_20260814"],
     "Movies behind the MAIN figure deck"),
    ("MOVIES_SUPPLEMENTAL", ["META_FIGURES_20260813_supplemental"],
     "Movies behind the SUPPLEMENTAL deck"),
    ("MOVIES_OTHER_THREE_FILES",
     ["NEW_FIGURES_20260804", "other", "NEW_TIMESTRIPS_20260804"],
     "Movies behind the other three Illustrator files"),
]

NAME_RE = re.compile(r"^(?P<deck>.+?)_AB(?P<ab>\d+)__(?P<fig>.+?)__(?P<batch>.+?)__(?P<role>ablation|monitoring)\.mp4$")


def parse(fn):
    m = NAME_RE.match(fn)
    if not m:
        return None
    d = m.groupdict()
    d["ab"] = int(d["ab"])
    d["file"] = fn
    return d


def poster_for(path, out_png):
    """A frame from about a third of the way in -- past any dark first frame, before anaphase."""
    if os.path.exists(out_png):
        return out_png
    cap = cv2.VideoCapture(path)
    if not cap.isOpened():
        return None
    n = int(cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0)
    cap.set(cv2.CAP_PROP_POS_FRAMES, max(0, n // 3))
    ok, fr = cap.read()
    if not ok:
        cap.set(cv2.CAP_PROP_POS_FRAMES, 0); ok, fr = cap.read()
    cap.release()
    if not ok:
        return None
    h, w = fr.shape[:2]
    s = min(1100.0 / w, 700.0 / h, 1.0)
    if s < 1:
        fr = cv2.resize(fr, (int(w * s), int(h * s)), interpolation=cv2.INTER_AREA)
    cv2.imwrite(out_png, fr)
    return out_png


def duration(path):
    cap = cv2.VideoCapture(path)
    if not cap.isOpened():
        return None
    n = cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0
    fps = cap.get(cv2.CAP_PROP_FPS) or 0
    cap.release()
    return (n / fps) if fps else None


def build():
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    os.makedirs(POSTER, exist_ok=True)
    movies = []
    for p in sorted(glob.glob(f"{MOV}/*.mp4")):
        d = parse(os.path.basename(p))
        if d:
            d["path"] = p
            movies.append(d)
        else:
            print(f"  unparsed filename, skipped: {os.path.basename(p)}")
    by_deck = collections.defaultdict(list)
    for m in movies:
        by_deck[m["deck"]].append(m)
    print(f"{len(movies)} movies across {len(by_deck)} decks: "
          + ", ".join(f"{k}={len(v)}" for k, v in sorted(by_deck.items())))

    made = []
    for fname, decks, subtitle in GROUPS:
        prs = Presentation()
        prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
        blank = prs.slide_layouts[6]

        def title_slide(text, sub, big=True):
            s = prs.slides.add_slide(blank)
            tb = s.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(11.7), Inches(2.4))
            tf = tb.text_frame; tf.word_wrap = True
            p0 = tf.paragraphs[0]; p0.text = text
            p0.runs[0].font.size = Pt(40 if big else 30); p0.runs[0].font.bold = True
            p1 = tf.add_paragraph(); p1.text = sub
            p1.runs[0].font.size = Pt(17); p1.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            return s

        total = sum(len(by_deck.get(d, [])) for d in decks)
        title_slide(subtitle, f"{total} processed movies  ·  built 2026-08-19 from "
                              f"3_MOVIES/timestrip_movies_20260817  ·  one slide per clip")

        for deck in decks:
            items = sorted(by_deck.get(deck, []), key=lambda m: (m["ab"], m["fig"], m["role"]))
            if not items:
                continue
            if len(decks) > 1:          # her instruction: a title slide leads off each file's group
                title_slide(DECK_TITLE.get(deck, deck), f"{len(items)} movies", big=False)
            for m in items:
                s = prs.slides.add_slide(blank)
                hb = s.shapes.add_textbox(Inches(0.45), Inches(0.25), Inches(12.4), Inches(0.9))
                htf = hb.text_frame; htf.word_wrap = True
                r = htf.paragraphs[0]; r.text = m["fig"]
                r.runs[0].font.size = Pt(22); r.runs[0].font.bold = True
                sub = htf.add_paragraph()
                dur = duration(m["path"])
                sub.text = (f"{DECK_TITLE.get(m['deck'], m['deck'])}  ·  artboard {m['ab']}  ·  "
                            f"{m['role']}  ·  cell {m['batch']}"
                            + (f"  ·  {dur:.0f} s" if dur else ""))
                sub.runs[0].font.size = Pt(12); sub.runs[0].font.color.rgb = RGBColor(0x44, 0x44, 0x44)

                png = poster_for(m["path"], f"{POSTER}/{m['file'][:-4]}.png")
                if png:
                    from PIL import Image
                    iw, ih = Image.open(png).size
                    maxw, maxh = Inches(9.6), Inches(5.0)
                    sc = min(maxw / iw, maxh / ih)
                    s.shapes.add_picture(png, Inches(0.45), Inches(1.35),
                                         width=int(iw * sc), height=int(ih * sc))
                fb = s.shapes.add_textbox(Inches(0.45), Inches(6.62), Inches(12.4), Inches(0.6))
                ftf = fb.text_frame; ftf.word_wrap = True
                fr_ = ftf.paragraphs[0]; fr_.text = m["path"]
                fr_.runs[0].font.size = Pt(9); fr_.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)

        out = f"{OUTDIR}/{fname}.pptx"
        prs.save(out)
        made.append((out, len(prs.slides._sldIdLst)))
        print(f"  wrote {out}  ({len(prs.slides._sldIdLst)} slides)")

    print("\nTo turn these into Google Slides: put the .pptx in Drive and open it with Google Slides,")
    print("or in an existing deck use File > Import slides.")
    return made


if __name__ == "__main__":
    os.makedirs(OUTDIR, exist_ok=True)
    build()
